"""Document parsing for uploads, plus the api-gateway helpers used by URL ingest."""

import io
import logging

import openpyxl
from docx import Document
from PyPDF2 import PdfReader

from core.config import get_config

logger = logging.getLogger(__name__)

_TEXT_MIMES = {"text/plain", "text/csv", "text/markdown"}
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _gateway_url(path: str) -> str:
    config = get_config()
    return f"{config.api_gateway_url.rstrip('/')}/{path.lstrip('/')}"


def _gateway_headers() -> dict:
    config = get_config()
    return {"X-API-Key": config.api_gateway_key} if config.api_gateway_key else {}


def parse_content(data: bytes, content_type: str, filename: str) -> str:
    """Parse raw bytes into plain text based on content type or filename extension."""
    name = filename.lower()

    if content_type in _TEXT_MIMES or name.endswith((".txt", ".md", ".markdown", ".csv")):
        return data.decode("utf-8", errors="replace")

    if content_type == "application/pdf" or name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(p for p in pages if p.strip())

    if content_type == _DOCX_MIME or name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)

    if content_type == _XLSX_MIME or name.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        sheets = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    rows.append("\t".join("" if c is None else str(c) for c in row))
            if rows:
                sheets.append(f"=== {sheet_name} ===\n" + "\n".join(rows))
        return "\n\n".join(sheets)

    logger.warning(f"Unknown content type '{content_type}' for '{filename}', attempting UTF-8 decode")
    return data.decode("utf-8", errors="replace")
